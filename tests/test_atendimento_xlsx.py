"""A planilha de atendimentos com link de WhatsApp (pedido de 21/09/2026).

O CSV entregava os números, mas quem ia usar a lista para falar com as pessoas
montava cada link na mão. A planilha entrega o link pronto — e, principalmente,
**um link que se refaz quando a mensagem muda**.

É fórmula, e não texto, por isso: a mensagem padrão mora numa célula da aba de
configuração, e a coluna WhatsApp aponta para ela. Se fossem links prontos, todo
ajuste de texto exigiria exportar de novo, e a planilha ficaria desatualizada em
silêncio.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta
from typing import Any

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.atendimento.infrastructure.models import Atendimento
from apps.customers.infrastructure.models import Customer
from apps.dashboards.exports.atendimento_xlsx import (
    MENSAGEM_PADRAO,
    montar_planilha,
)
from apps.shared.context import set_current_organization
from apps.tenancy.models import Organization


def _linha(**over: Any) -> dict[str, Any]:
    base = {
        "customer_name": "Fulano de Tal",
        "customer_document": "12345678901",
        "telefone_conversa": "(15) 99150-1696",
        "telefone_cadastro": "(15) 98825-1772",
        "opened_at_str": "19/09 14:30",
        "atendente_nome": "José",
        "departamento_nome": "Suporte",
        "categorias": ["Lentidão"],
        "protocol": "P1",
        "status_label": "Finalizado",
    }
    base.update(over)
    return base


def _abre(wb: Any) -> Any:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return load_workbook(buffer)


class TestPlanilha:
    def test_tem_as_duas_abas_e_a_de_dados_vem_na_frente(self) -> None:
        """A configuração é ajuste, não destino: quem abre quer ver a lista."""
        wb = _abre(montar_planilha(
            [_linha()], exportado_em=datetime(2026, 9, 21, 10, 30), recorte="01/09 a 17/09"
        ))
        assert wb.sheetnames == ["Configuração", "Atendimentos"]
        assert wb.active.title == "Atendimentos"

    def test_link_e_formula_que_aponta_para_a_mensagem(self) -> None:
        """O ponto da planilha: editar a mensagem numa célula muda todos os
        links. Link pronto obrigaria a exportar de novo a cada ajuste."""
        wb = _abre(montar_planilha(
            [_linha()], exportado_em=datetime(2026, 9, 21, 10, 30), recorte="x"
        ))
        formula = wb["Atendimentos"]["K2"].value
        assert formula.startswith("=IF(")
        assert "Configuração!$B$2" in formula
        assert "https://wa.me/" in formula
        assert "HYPERLINK" in formula

    def test_link_usa_o_numero_da_conversa_sem_formatacao(self) -> None:
        """wa.me não aceita parênteses nem espaço — a fórmula limpa o número."""
        wb = _abre(montar_planilha(
            [_linha()], exportado_em=datetime(2026, 9, 21, 10, 30), recorte="x"
        ))
        formula = wb["Atendimentos"]["K2"].value
        assert 'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(C2," ",""),"(",""),")","")' in formula

    def test_nome_do_cliente_entra_na_mensagem(self) -> None:
        wb = _abre(montar_planilha(
            [_linha()], exportado_em=datetime(2026, 9, 21, 10, 30), recorte="x"
        ))
        assert '"{nome}",A2' in wb["Atendimentos"]["K2"].value

    def test_sem_numero_nao_gera_link(self) -> None:
        """`wa.me/` sem destinatário abre o WhatsApp em branco — pior que célula
        vazia, porque parece que funcionou."""
        wb = _abre(montar_planilha(
            [_linha(telefone_conversa="")],
            exportado_em=datetime(2026, 9, 21, 10, 30), recorte="x",
        ))
        assert wb["Atendimentos"]["K2"].value.startswith('=IF(C2="","",')

    def test_configuracao_carrega_o_contexto_da_exportacao(self) -> None:
        """A planilha perde o contexto no primeiro encaminhamento de e-mail —
        sem data e recorte escritos dentro dela, ninguém sabe se está olhando
        ontem ou três meses atrás."""
        wb = _abre(montar_planilha(
            [_linha(), _linha(telefone_conversa="")],
            exportado_em=datetime(2026, 9, 21, 10, 30),
            recorte="01/09 a 17/09 · Suporte",
        ))
        valores = {
            wb["Configuração"][f"A{i}"].value: wb["Configuração"][f"B{i}"].value
            for i in range(1, 12)
        }
        assert valores["Exportado em"] == "21/09/2026 10:30"
        assert valores["Recorte"] == "01/09 a 17/09 · Suporte"
        assert valores["Atendimentos"] == 2
        assert valores["Com número para WhatsApp"] == 1

    def test_mensagem_personalizada_substitui_a_padrao(self) -> None:
        wb = _abre(montar_planilha(
            [_linha()], exportado_em=datetime(2026, 9, 21, 10, 30), recorte="x",
            mensagem="Oi {nome}, tudo certo?",
        ))
        assert wb["Configuração"]["B2"].value == "Oi {nome}, tudo certo?"

    def test_mensagem_padrao_tem_o_marcador_de_nome(self) -> None:
        assert "{nome}" in MENSAGEM_PADRAO


@pytest.mark.django_db
@pytest.mark.filterwarnings("ignore:No directory at:UserWarning")
class TestExportPelaTela:
    def test_format_xlsx_devolve_planilha(
        self, client: Any, user_a: Any, organization_a: Organization
    ) -> None:
        set_current_organization(organization_a)
        cliente = Customer.objects.create(
            organization=organization_a, source_type="IXC", external_id="c1",
            name="Fulano", phone="(15) 98825-1772",
        )
        Atendimento.objects.create(
            organization=organization_a, source_type="OPA", external_id="at-1",
            customer=cliente, customer_external_id="c1", customer_name="Fulano",
            status=Atendimento.Status.CLOSED, canal="whatsapp",
            opened_at=timezone.now() - timedelta(hours=2),
            raw_extras={"canal_cliente": "5515991501696@c.us"},
        )
        client.force_login(user_a)
        hoje = timezone.now().date()
        url = (
            f"/operations/atendimento-lista/?de={hoje - timedelta(days=1)}"
            f"&ate={hoje}&foco=todos&origem=atendimento&format=xlsx"
        )
        resp = client.get(url)
        assert resp.status_code == 200
        assert "spreadsheetml" in resp["Content-Type"]
        assert resp["Content-Disposition"].endswith('.xlsx"')

        wb = load_workbook(io.BytesIO(resp.content))
        assert wb["Atendimentos"]["C2"].value == "(15) 99150-1696"
        assert "wa.me" in wb["Atendimentos"]["K2"].value

    def test_csv_continua_funcionando(
        self, client: Any, user_a: Any, organization_a: Organization
    ) -> None:
        """Quem já usa o CSV num fluxo não pode ser quebrado por uma melhoria de
        formato."""
        client.force_login(user_a)
        hoje = timezone.now().date()
        resp = client.get(
            f"/operations/atendimento-lista/?de={hoje - timedelta(days=1)}"
            f"&ate={hoje}&foco=todos&origem=atendimento&format=csv"
        )
        assert resp.status_code == 200
        assert resp["Content-Type"].startswith("text/csv")
